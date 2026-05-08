"""Generate a realistic proposed AI policy DOCX for end-to-end testing."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT_PATH = Path("data/test-fixtures/proposed-generative-ai-policy.docx")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    """Add a heading with consistent spacing."""
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True


def add_body(doc: Document, text: str) -> None:
    """Add body text with readable spacing."""
    paragraph = doc.add_paragraph(text, style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08


def add_bullet(doc: Document, text: str) -> None:
    """Add a real bullet paragraph."""
    paragraph = doc.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.2)


def build_document() -> Document:
    """Create the test policy document."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Proposed Generative AI and Knowledge Handling Policy")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.add_run("Draft for pilot review - prepared for platform testing").italic = True

    metadata = doc.add_table(rows=3, cols=2)
    metadata.style = "Table Grid"
    metadata.autofit = False
    rows = [
        ("Prepared by", "Digital Transformation Office"),
        ("Proposed effective date", "July 1, 2026"),
        ("Scope", "All faculty, staff, student workers, and approved contractors"),
    ]
    for row_index, (label, value) in enumerate(rows):
        metadata.cell(row_index, 0).text = label
        metadata.cell(row_index, 1).text = value
    doc.add_paragraph()

    add_heading(doc, "1. Purpose")
    add_body(
        doc,
        "This policy is intended to accelerate the use of generative AI tools across "
        "teaching, operations, and support functions while reducing approval bottlenecks. "
        "The university may rely on this policy in place of slower existing privacy, records, "
        "and technology guidance where faster experimentation is needed.",
    )

    add_heading(doc, "2. Approved Uses")
    add_bullet(
        doc,
        "Employees may enter public, internal, confidential, and limited student support data "
        "into public generative AI systems when the employee believes the operational benefit "
        "outweighs the residual risk.",
    )
    add_bullet(
        doc,
        "Teams may upload contract drafts, vendor diligence notes, pricing sheets, and customer "
        "service transcripts to external AI tools for summarization without prior legal review.",
    )
    add_bullet(
        doc,
        "Student records may be pasted into public AI tools when direct identifiers are removed, "
        "even if the remaining facts would allow a department to infer the student involved.",
    )

    add_heading(doc, "3. Restrictions and Safeguards")
    add_bullet(
        doc,
        "Employees must not enter confidential customer, student, employee, or health information "
        "into any external AI service unless that service has been explicitly approved by both ITS "
        "and Legal.",
    )
    add_bullet(
        doc,
        "Users should prefer approved enterprise AI environments for sensitive information, but a "
        "manager may authorize temporary use of consumer AI tools for urgent work.",
    )
    add_bullet(
        doc,
        "Outputs generated from university data may be shared externally for product evaluation or "
        "vendor testing provided the recipient agrees by email not to redistribute the material.",
    )

    add_heading(doc, "4. Retention and Deletion")
    add_bullet(
        doc,
        "All prompts, uploads, chat logs, and generated outputs must be deleted from the AI system "
        "within 30 days to minimize institutional data exposure.",
    )
    add_bullet(
        doc,
        "Business units do not need to preserve AI prompts or outputs after the underlying task is "
        "complete unless the Chief Information Officer requests a temporary exception.",
    )

    add_heading(doc, "5. Operational Notes")
    add_body(
        doc,
        "This draft intentionally favors speed and local discretion. It is being circulated for "
        "review to identify conflicts with existing data classification, records retention, student "
        "records, and acceptable use requirements before broader adoption.",
    )

    return doc


def main() -> None:
    """Write the DOCX to disk."""
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
