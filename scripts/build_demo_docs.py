"""Convert demo markdown policies to .docx files for ingestion.

Run once before the demo to materialize the .docx versions of the markdown
sources under data/demo/. Idempotent — safe to re-run.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


_DATA_DEMO = Path(__file__).resolve().parent.parent / "data" / "demo"
_SOURCES = [
    "acme_existing_ai_policy.md",
    "acme_proposed_ai_policy.md",
]


def _add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def _convert(md_path: Path, docx_path: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    for line in lines:
        stripped = line.rstrip()
        if not stripped:
            document.add_paragraph()
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        else:
            _add_paragraph(document, stripped)

    document.save(str(docx_path))
    print(f"wrote {docx_path}")


def main() -> None:
    if not _DATA_DEMO.exists():
        raise SystemExit(f"Demo directory not found: {_DATA_DEMO}")
    for source in _SOURCES:
        md_path = _DATA_DEMO / source
        if not md_path.exists():
            raise SystemExit(f"Missing markdown source: {md_path}")
        docx_path = md_path.with_suffix(".docx")
        _convert(md_path, docx_path)


if __name__ == "__main__":
    main()
