"""Tests for src/ingestion/converter.py — DocumentConverter format handling."""

import pytest
from PIL import Image
from src.ingestion.converter import DocumentConverter


class TestImageConversion:
    def test_convert_png(self, sample_image):
        """PNG image should return a list with one PIL Image."""
        result = DocumentConverter.convert(sample_image)
        assert isinstance(result, list)
        assert len(result) == 1
        assert isinstance(result[0], Image.Image)

    def test_convert_jpg(self, tmp_path):
        """JPEG image should return a list with one PIL Image."""
        img = Image.new("RGB", (10, 10), "green")
        jpg_path = tmp_path / "test.jpg"
        img.save(jpg_path)

        result = DocumentConverter.convert(jpg_path)
        assert len(result) == 1
        assert isinstance(result[0], Image.Image)


class TestPdfConversion:
    def test_convert_pdf(self, sample_pdf):
        """PDF should return a list of PIL Images (one per page)."""
        result = DocumentConverter.convert(sample_pdf)
        assert isinstance(result, list)
        assert len(result) >= 1
        assert isinstance(result[0], Image.Image)


class TestOfficeConversion:
    def test_convert_docx(self, sample_docx):
        """DOCX should return a list of strings."""
        result = DocumentConverter.convert(sample_docx)
        assert isinstance(result, list)
        assert len(result) >= 1
        assert isinstance(result[0], str)
        assert "Test paragraph" in result[0]

    def test_convert_docx_includes_tables(self, tmp_path):
        """DOCX extraction should preserve document structure as markdown."""
        import docx

        doc = docx.Document()
        doc.add_heading("Policy introduction", level=1)
        doc.add_paragraph("Employees may use approved tools.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Control"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Data review"
        table.cell(1, 1).text = "Compliance"
        path = tmp_path / "table.docx"
        doc.save(path)

        result = DocumentConverter.convert(path)

        assert "# Policy introduction" in result[0]
        assert "Employees may use approved tools." in result[0]
        assert "| Control | Owner |" in result[0]
        assert "| --- | --- |" in result[0]
        assert "| Data review | Compliance |" in result[0]

    def test_convert_pptx(self, sample_pptx):
        """PPTX should return a list of strings (one per slide)."""
        result = DocumentConverter.convert(sample_pptx)
        assert isinstance(result, list)
        assert len(result) >= 1
        assert isinstance(result[0], str)
        assert "Test Slide Title" in result[0]

    def test_convert_xlsx(self, sample_xlsx):
        """XLSX should return markdown tables (one per sheet)."""
        result = DocumentConverter.convert(sample_xlsx)
        assert isinstance(result, list)
        assert len(result) >= 1
        assert isinstance(result[0], str)
        assert "## TestSheet" in result[0]
        assert "| Header1 | Header2 |" in result[0]
        assert "| --- | --- |" in result[0]


class TestEdgeCases:
    def test_unsupported_format(self, tmp_path):
        """Unsupported file extension should raise ValueError."""
        bad_file = tmp_path / "test.xyz"
        bad_file.write_text("content")
        with pytest.raises(ValueError, match="Unsupported"):
            DocumentConverter.convert(bad_file)

    def test_file_not_found(self):
        """Non-existent file should raise FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            DocumentConverter.convert("/nonexistent/path/file.pdf")

    def test_supported_extensions(self):
        """get_supported_extensions should return expected list."""
        exts = DocumentConverter.get_supported_extensions()
        assert ".pdf" in exts
        assert ".png" in exts
        assert ".docx" in exts
        assert ".xlsx" in exts
        assert ".pptx" in exts
