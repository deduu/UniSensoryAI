"""Shared test fixtures for UniSensoryAI test suite."""

import pytest
from unittest.mock import AsyncMock, MagicMock
from datetime import datetime, timezone
from PIL import Image

from src.db.models.ingestion import IngestionJob, IngestionStatus
from src.db.models.document import Document


@pytest.fixture
def mock_session():
    """AsyncMock of SQLAlchemy AsyncSession."""
    session = AsyncMock()
    session.commit = AsyncMock()
    session.flush = AsyncMock()
    session.refresh = AsyncMock()
    session.delete = AsyncMock()
    session.add = MagicMock()
    return session


@pytest.fixture
def mock_session_manager(mock_session):
    """MagicMock wrapping mock_session as async context manager."""
    manager = MagicMock()
    manager.session.return_value.__aenter__ = AsyncMock(return_value=mock_session)
    manager.session.return_value.__aexit__ = AsyncMock(return_value=False)
    return manager


@pytest.fixture
def sample_ingestion_job():
    """MagicMock of IngestionJob with sensible defaults."""
    job = MagicMock(spec=IngestionJob)
    job.id = "test-job-001"
    job.status = IngestionStatus.PENDING
    job.file_path = "/tmp/test.pdf"
    job.meta = {"model_id": "qwen2-vl-2b-instruct"}
    job.created_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
    job.updated_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
    job.documents = []
    return job


@pytest.fixture
def sample_document():
    """MagicMock of Document with sensible defaults."""
    doc = MagicMock(spec=Document)
    doc.id = "test-doc-001"
    doc.job_id = "test-job-001"
    doc.filename = "test.pdf"
    doc.content_text = "Extracted text content"
    doc.content_chunks = [{"page": 1, "text": "chunk"}]
    doc.created_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
    return doc


@pytest.fixture
def tmp_upload_dir(tmp_path):
    """Temporary upload directory."""
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    return upload_dir


@pytest.fixture
def sample_image(tmp_path):
    """Small 10x10 PNG image file in tmp_path. Returns path."""
    img = Image.new("RGB", (10, 10), color="red")
    img_path = tmp_path / "test_image.png"
    img.save(img_path)
    return img_path


@pytest.fixture
def sample_xlsx(tmp_path):
    """Minimal Excel file in tmp_path. Returns path."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSheet"
    ws["A1"] = "Header1"
    ws["B1"] = "Header2"
    ws["A2"] = "Value1"
    ws["B2"] = "Value2"
    xlsx_path = tmp_path / "test.xlsx"
    wb.save(xlsx_path)
    return xlsx_path


@pytest.fixture
def sample_docx(tmp_path):
    """Minimal Word document in tmp_path. Returns path."""
    import docx

    doc = docx.Document()
    doc.add_paragraph("Test paragraph one.")
    doc.add_paragraph("Test paragraph two.")
    docx_path = tmp_path / "test.docx"
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def sample_pptx(tmp_path):
    """Minimal PowerPoint file in tmp_path. Returns path."""
    import pptx

    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Test Slide Title"
    pptx_path = tmp_path / "test.pptx"
    prs.save(pptx_path)
    return pptx_path


@pytest.fixture
def sample_pdf(tmp_path):
    """Minimal single-page PDF in tmp_path. Returns path."""
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=200, height=200)
    page.insert_text((50, 100), "Test PDF content")
    pdf_path = tmp_path / "test.pdf"
    doc.save(pdf_path)
    doc.close()
    return pdf_path
